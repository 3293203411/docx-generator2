import re
from docx.enum.text import WD_ALIGN_PARAGRAPH

def is_markdown_table(text):
    """判断是否为标准 Markdown 表格格式"""
    if not text or '|' not in text:
        return False
    lines = [l.strip() for l in text.strip().splitlines() if l.strip()]
    if len(lines) < 3:
        return False
    # 检查第二行是否为分隔符行（仅允许 |, -, :, 空格）
    sep_line = lines[1].strip('|')
    return bool(re.match(r'^[\s|:-]+$', sep_line))

def parse_markdown_table(text):
    """解析 Markdown 表格，返回 (headers, alignments, data_rows)"""
    lines = [l.strip() for l in text.strip().splitlines() if l.strip()]
    
    def split_row(line):
        return [cell.strip() for cell in line.strip('|').split('|')]

    headers = split_row(lines[0])
    sep_parts = split_row(lines[1])
    
    # 提取每列对齐方式
    alignments = []
    for part in sep_parts:
        part = part.strip()
        if part.startswith(':') and part.endswith(':'):
            alignments.append(WD_ALIGN_PARAGRAPH.CENTER)
        elif part.endswith(':'):
            alignments.append(WD_ALIGN_PARAGRAPH.RIGHT)
        else:
            alignments.append(WD_ALIGN_PARAGRAPH.LEFT)
            
    data_rows = [split_row(line) for line in lines[2:]]
    return headers, alignments, data_rows

def replace_placeholder_with_auto_table(doc, placeholder, table_text):
    """将占位符段落替换为 Markdown 解析后的 Word 表格"""
    target_para = None
    body = doc.element.body
    for element in body:
        if isinstance(element, CT_P):
            para = Paragraph(element, doc)
            if placeholder in "".join(run.text for run in para.runs):
                target_para = para
                break
    if target_para is None:
        return False

    headers, alignments, data_rows = parse_markdown_table(table_text)
    if not headers or not data_rows:
        return False

    cols_count = len(headers)
    rows_count = len(data_rows) + 1  # +1 表头
    table = doc.add_table(rows=rows_count, cols=cols_count)
    table.style = 'Table Grid'

    # 1. 填充表头
    for c_idx, header_text in enumerate(headers):
        cell = table.cell(0, c_idx)
        set_cell_border(cell)
        set_cell_background(cell, 'E6E6E6')
        align = alignments[c_idx] if c_idx < len(alignments) else WD_ALIGN_PARAGRAPH.LEFT
        set_cell_text(cell, header_text, bold=True, align=align)

    # 2. 填充数据行
    for r_idx, row_data in enumerate(data_rows, start=1):
        # 补齐或截断单元格数量，防止越界
        row_data = row_data[:cols_count]
        row_data += [''] * (cols_count - len(row_data))
        
        for c_idx, cell_text in enumerate(row_data):
            cell = table.cell(r_idx, c_idx)
            set_cell_border(cell)
            if r_idx % 2 == 0:  # 偶数行斑马纹
                set_cell_background(cell, 'F5F5F5')
            align = alignments[c_idx] if c_idx < len(alignments) else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cell, cell_text, bold=False, align=align)

    # 3. 插入表格并移除原占位符段落
    target_para._element.addprevious(table._tbl)
    parent = target_para._element.getparent()
    parent.remove(target_para._element)
    return True

def process_document(doc, keys, values):
    # ================= 自动识别并转换任意 Markdown 表格 =================
    new_keys, new_values = [], []
    for k, v in zip(keys, values):
        # 如果值符合 Markdown 表格特征，则自动转换，不再作为普通文本处理
        if is_markdown_table(v):
            placeholder = f"{{{{{k}}}}}"
            if replace_placeholder_with_auto_table(doc, placeholder, v):
                continue  # 转换成功，跳过
        new_keys.append(k)
        new_values.append(v)
    keys, values = new_keys, new_values

    # ================= 研发内容表格 (JSON格式，保留原逻辑) =================
    table_placeholder = '研发内容完成情况'
    table_data_key_idx = next((i for i, k in enumerate(keys) if k == table_placeholder), None)
    if table_data_key_idx is not None:
        try:
            table_data = json.loads(values[table_data_key_idx])
            if isinstance(table_data, list):
                replace_placeholder_with_table(doc, f'{{{{{table_placeholder}}}}}', table_data)
                keys.pop(table_data_key_idx)
                values.pop(table_data_key_idx)
        except Exception:
            pass

    # ================= 普通文本替换 =================
    for element in doc.element.body:
        if isinstance(element, CT_P):
            process_paragraph(Paragraph(element, doc), keys, values)
    for table in doc.tables:
        process_table(table, keys, values)
